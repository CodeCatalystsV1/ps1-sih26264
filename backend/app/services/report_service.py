import io
import csv
from datetime import datetime, timezone, timedelta
from typing import List, Dict, Any, Optional
from sqlalchemy.orm import Session
from sqlalchemy import func, distinct
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from backend.app.models.entities import (
    Complaint, Case, Institution, Bureau, Region, OverlapFlag, Transfer, WorkflowHistory
)
from backend.app.schemas.all_schemas import DashboardMetricsOut

def get_dashboard_metrics(db: Session, bureau_id: Optional[int] = None, region_id: Optional[int] = None) -> DashboardMetricsOut:
    now = datetime.now(timezone.utc)
    
    # Base queries
    comp_q = db.query(Complaint)
    case_q = db.query(Case)

    if bureau_id:
        case_q = case_q.filter(Case.current_bureau_id == bureau_id)
        comp_q = comp_q.filter(Complaint.assigned_bureau_id == bureau_id)
    if region_id:
        case_q = case_q.join(Institution).filter(Institution.region_id == region_id)
        comp_q = comp_q.join(Institution).filter(Institution.region_id == region_id)

    total_complaints = comp_q.count()

    # Case statuses
    pending_statuses = ["Received", "Acknowledged", "Under Scrutiny", "Assigned"]
    under_process_statuses = ["Institution Notice Issued", "Awaiting Response", "Under Verification", "Committee Review", "Action Taken"]
    disposed_statuses = ["Disposed", "Closed – Action Taken", "Closed – No Action Required"]

    pending_cases = case_q.filter(Case.status.in_(pending_statuses)).count()
    cases_under_process = case_q.filter(Case.status.in_(under_process_statuses)).count()
    disposed_cases = case_q.filter(Case.status.in_(disposed_statuses)).count()
    transferred_cases = case_q.filter(Case.status == "Transferred").count()

    overdue_cases = case_q.filter(
        Case.status.notin_(disposed_statuses),
        Case.applicable_deadline_at.isnot(None),
        Case.applicable_deadline_at < now
    ).count()

    overlap_q = db.query(OverlapFlag).filter(
        OverlapFlag.relationship_state.in_(["POTENTIAL_OVERLAP", "UNDER_REVIEW", "COORDINATION_REQUESTED", "JOINT_ACTION_IN_PROGRESS", "ESCALATED"])
    )
    if region_id:
        overlap_q = overlap_q.join(Institution).filter(Institution.region_id == region_id)
    potential_overlaps = overlap_q.count()

    active_inst_q = db.query(distinct(Case.institution_id)).filter(Case.status.notin_(disposed_statuses))
    if bureau_id:
        active_inst_q = active_inst_q.filter(Case.current_bureau_id == bureau_id)
    if region_id:
        active_inst_q = active_inst_q.join(Institution).filter(Institution.region_id == region_id)
    institutions_with_active_cases = active_inst_q.count()

    return DashboardMetricsOut(
        total_complaints=total_complaints,
        pending_cases=pending_cases,
        cases_under_process=cases_under_process,
        disposed_cases=disposed_cases,
        transferred_cases=transferred_cases,
        overdue_cases=overdue_cases,
        potential_overlaps=potential_overlaps,
        institutions_with_active_cases=institutions_with_active_cases
    )

def get_ageing_distribution(db: Session, bureau_id: Optional[int] = None) -> List[Dict[str, Any]]:
    now = datetime.now(timezone.utc)
    cases = db.query(Case).filter(
        Case.status.notin_(["Disposed", "Closed – Action Taken", "Closed – No Action Required"])
    )
    if bureau_id:
        cases = cases.filter(Case.current_bureau_id == bureau_id)
    cases = cases.all()

    buckets = {"0-7 days": 0, "8-30 days": 0, "31-60 days": 0, "60+ days": 0}
    for c in cases:
        rec_date = c.received_at
        if rec_date.tzinfo is None:
            rec_date = rec_date.replace(tzinfo=timezone.utc)
        age = (now - rec_date).days
        if age <= 7:
            buckets["0-7 days"] += 1
        elif age <= 30:
            buckets["8-30 days"] += 1
        elif age <= 60:
            buckets["31-60 days"] += 1
        else:
            buckets["60+ days"] += 1

    return [{"band": k, "count": v} for k, v in buckets.items()]

def get_bureau_workload_report(db: Session) -> List[Dict[str, Any]]:
    bureaus = db.query(Bureau).filter(Bureau.active == True).all()
    results = []
    now = datetime.now(timezone.utc)
    for b in bureaus:
        total_assigned = db.query(Case).filter(Case.current_bureau_id == b.id).count()
        pending = db.query(Case).filter(Case.current_bureau_id == b.id, Case.status.in_(["Received", "Acknowledged", "Under Scrutiny", "Assigned"])).count()
        under_process = db.query(Case).filter(Case.current_bureau_id == b.id, Case.status.in_(["Institution Notice Issued", "Awaiting Response", "Under Verification", "Committee Review", "Action Taken"])).count()
        disposed = db.query(Case).filter(Case.current_bureau_id == b.id, Case.status.in_(["Disposed", "Closed – Action Taken", "Closed – No Action Required"])).count()
        overdue = db.query(Case).filter(
            Case.current_bureau_id == b.id,
            Case.status.notin_(["Disposed", "Closed – Action Taken", "Closed – No Action Required"]),
            Case.applicable_deadline_at.isnot(None),
            Case.applicable_deadline_at < now
        ).count()

        results.append({
            "bureau_id": b.id,
            "bureau_name": b.name,
            "bureau_code": b.code,
            "total_assigned": total_assigned,
            "pending_count": pending,
            "under_process_count": under_process,
            "disposed_count": disposed,
            "overdue_count": overdue
        })
    return results

def get_regional_distribution_report(db: Session) -> List[Dict[str, Any]]:
    regions = db.query(Region).filter(Region.active == True).all()
    results = []
    now = datetime.now(timezone.utc)
    for r in regions:
        inst_count = db.query(Institution).filter(Institution.region_id == r.id).count()
        comp_count = db.query(Complaint).join(Institution).filter(Institution.region_id == r.id).count()
        case_count = db.query(Case).join(Institution).filter(Institution.region_id == r.id).count()
        pending = db.query(Case).join(Institution).filter(
            Institution.region_id == r.id,
            Case.status.in_(["Received", "Acknowledged", "Under Scrutiny", "Assigned"])
        ).count()
        disposed = db.query(Case).join(Institution).filter(
            Institution.region_id == r.id,
            Case.status.in_(["Disposed", "Closed – Action Taken", "Closed – No Action Required"])
        ).count()
        overdue = db.query(Case).join(Institution).filter(
            Institution.region_id == r.id,
            Case.status.notin_(["Disposed", "Closed – Action Taken", "Closed – No Action Required"]),
            Case.applicable_deadline_at.isnot(None),
            Case.applicable_deadline_at < now
        ).count()

        results.append({
            "region_id": r.id,
            "region_name": r.name,
            "region_code": r.code,
            "institution_count": inst_count,
            "complaint_count": comp_count,
            "case_count": case_count,
            "pending_count": pending,
            "disposed_count": disposed,
            "overdue_count": overdue
        })
    return results

# ---------------- EXPORT IMPLEMENTATIONS ----------------
def export_cases_to_csv(cases: List[Case]) -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Case Reference", "Institution Name", "Permanent ID", "Current Bureau",
        "Assigned Officer", "Status", "Priority", "Received Date",
        "Last Action", "Next Action", "Deadline"
    ])
    for c in cases:
        writer.writerow([
            c.case_reference,
            c.institution.institution_name if c.institution else "",
            c.institution.permanent_id if c.institution else "",
            c.current_bureau.name if c.current_bureau else "",
            c.current_officer.display_name if c.current_officer else "Unassigned",
            c.status,
            c.priority,
            c.received_at.strftime("%Y-%m-%d %H:%M"),
            c.last_action or "",
            c.next_action or "",
            c.applicable_deadline_at.strftime("%Y-%m-%d") if c.applicable_deadline_at else "Deadline not configured"
        ])
    return output.getvalue()

def export_cases_to_excel(cases: List[Case]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AICTE ICMS Cases"

    headers = [
        "Case Reference", "Institution Name", "Permanent ID", "Bureau",
        "Assigned Officer", "Status", "Priority", "Received Date",
        "Last Action", "Next Action", "Deadline"
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="17365D", end_color="17365D", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for c in cases:
        ws.append([
            c.case_reference,
            c.institution.institution_name if c.institution else "",
            c.institution.permanent_id if c.institution else "",
            c.current_bureau.name if c.current_bureau else "",
            c.current_officer.display_name if c.current_officer else "Unassigned",
            c.status,
            c.priority,
            c.received_at.strftime("%Y-%m-%d %H:%M"),
            c.last_action or "",
            c.next_action or "",
            c.applicable_deadline_at.strftime("%Y-%m-%d") if c.applicable_deadline_at else "Deadline not configured"
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def export_cases_to_pdf(cases: List[Case], report_title: str = "AICTE ICMS — Official Case Records Report") -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        rightMargin=20, leftMargin=20, topMargin=25, bottomMargin=25
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=14,
        textColor=colors.HexColor('#17365D'),
        spaceAfter=6
    )
    meta_style = ParagraphStyle(
        'Meta',
        parent=styles['Normal'],
        fontSize=9,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=12
    )
    cell_style = ParagraphStyle(
        'Cell',
        parent=styles['Normal'],
        fontSize=8,
        leading=10
    )

    elements = [
        Paragraph(report_title, title_style),
        Paragraph(f"Generated on {datetime.now(timezone.utc).strftime('%d-%b-%Y %H:%M UTC')} | Confidential Administrative Record", meta_style),
        Spacer(1, 8)
    ]

    data = [[
        Paragraph("<b>Case Ref</b>", cell_style),
        Paragraph("<b>Institution / Perm ID</b>", cell_style),
        Paragraph("<b>Bureau</b>", cell_style),
        Paragraph("<b>Officer</b>", cell_style),
        Paragraph("<b>Status</b>", cell_style),
        Paragraph("<b>Priority</b>", cell_style),
        Paragraph("<b>Received</b>", cell_style),
        Paragraph("<b>Deadline</b>", cell_style)
    ]]

    for c in cases[:150]: # Cap per page bundle
        inst_text = f"{c.institution.institution_name[:35]}<br/>({c.institution.permanent_id})" if c.institution else "N/A"
        officer = c.current_officer.display_name if c.current_officer else "Unassigned"
        deadline = c.applicable_deadline_at.strftime("%d-%b-%Y") if c.applicable_deadline_at else "Not configured"

        data.append([
            Paragraph(c.case_reference, cell_style),
            Paragraph(inst_text, cell_style),
            Paragraph(c.current_bureau.code if c.current_bureau else "", cell_style),
            Paragraph(officer, cell_style),
            Paragraph(c.status, cell_style),
            Paragraph(c.priority, cell_style),
            Paragraph(c.received_at.strftime("%d-%b-%Y"), cell_style),
            Paragraph(deadline, cell_style)
        ])

    table = Table(data, colWidths=[80, 200, 70, 90, 100, 55, 75, 90])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#17365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D9E0E8')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
    ]))

    elements.append(table)
    doc.build(elements)
    return buf.getvalue()
